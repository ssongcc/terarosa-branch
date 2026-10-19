# -*- coding: utf-8 -*-
"""
지점 원두 발주서 → 우체국 택배 접수파일 변환 로직 (앱/스크립트 공용)

규칙
  1) [어센틱에스프레소블렌드 1kg] : 지점별 수량(=kg) → 15kg 기준 박스 분할(올림), 박스별 고르게 분배
  2) 그 외 품명 뒤 중량(g/kg)이 붙은 원두 : 지점별 중량 합산 → 15kg 기준 박스 분할(올림)
  3) 1박스 = 1행, 수취인명 뒤 (박스번호/총박스)  예) (1/3) (2/3) (3/3), 1박스면 (1/1)
  4) 어센틱 박스 품목(F열) = '어센틱에스프레소블렌드1kg 10kg' 처럼 박스별 중량 표기
  5) 강릉 지점(본점·경포·임당·사천·신라모노), 주소파일에 없는 지점은 제외
"""
import io, re, math, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BOX_KG = 15
AUTHENTIC = "어센틱에스프레소블렌드1kg"           # 공백 제거 후 비교
AUTHENTIC_LABEL = "어센틱에스프레소블렌드1kg"     # 송장 품목(F열) 표기
EXCLUDE = {"본점", "경포", "임당", "사천", "신라모노"}
# 발주서 헤더 → 주소파일 수취인명에 포함된 키워드 (헤더와 다를 때만)
ALIAS = {
    "서종": "서종점", "예당": "예술의전당", "양재": "양재역", "포스코": "포스코센터",
    "현대": "국립현대미술관", "중문": "중문", "청계": "청계광장", "구의": "구의역",
    "잠실한강": "잠실한강공원", "역삼": "역삼역", "용산": "용산아이파크",
    "스타필드하남": "스타필드하남", "청주": "청주지웰", "교육서비스": "교육서비스",
    "IFC": "IFC몰", "현대송도": "현대아울렛송도", "현대동대문": "현대아울렛동대문",
    "성수": "성수", "신세계대구": "신세계대구", "청량리": "청량리",
}
WEIGHT_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(kg|g)", re.I)
HEADERS = ["수취인명", "우편번호", "상세주소", "연락처 1 일반전화", "연락처 2 휴대폰", "품목", "메모"]


def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace("\xa0", " ").replace("\n", " ").strip()


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def boxes(kg):
    kg = round(kg, 6)
    return math.ceil(kg / BOX_KG - 1e-9) if kg > 0 else 0


def split_even(kg, n):
    """kg를 n박스에 고르게 분배. 정수 kg면 정수로, 남는 1kg은 앞 박스부터"""
    if n <= 0:
        return []
    kg = round(kg, 6)
    if float(kg).is_integer():
        base, rem = divmod(int(kg), n)
        return [base + (1 if i < rem else 0) for i in range(n)]
    part = round(kg / n, 2)
    return [part] * (n - 1) + [round(kg - part * (n - 1), 2)]


ADDR_KEYS = ["name", "zip", "addr", "tel1", "tel2", "item", "memo"]


def load_addresses(src):
    """src: 경로 / file-like(xlsx) / 주소 리스트(Secrets의 [[addresses]])"""
    if isinstance(src, (list, tuple)):
        rows = []
        for a in src:
            row = [s(a.get(k, "")) for k in ADDR_KEYS] if hasattr(a, "get") else [s(x) for x in list(a)[:7]]
            if row[0]:
                rows.append(row + [""] * (7 - len(row)))
        return rows
    ws = openpyxl.load_workbook(src, data_only=True).worksheets[0]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not s(r[0]):
            continue
        rows.append([s(r[i]) if i < len(r) else "" for i in range(7)])
    return rows


def match_address(header, addrs):
    key = ALIAS.get(header, header)
    hits = [a for a in addrs if key in a[0].replace("[[원두]테라로사", "")]
    if len(hits) == 1:
        return hits[0], ""
    return None, ("주소 미등록 → 제외" if not hits else "주소 중복매칭 → 확인필요")


def allocate_other(items, n):
    """
    기타 원두를 n박스에 배분 (박스당 15kg 이하).
    items: [(품명, 단위kg, 수량)] → [[(품명, 수량, kg), ...], ...]
    같은 품목은 가능한 한 한 박스에 통째로 넣고, 넘칠 때만 나눠 담음. 자리가 없으면 박스 추가.
    """
    cap = int(BOX_KG * 1000)
    bx = [dict() for _ in range(max(n, 1))]
    load = [0] * len(bx)
    order = sorted(items, key=lambda x: -round(x[1] * 1000) * x[2])
    for name, w, q in order:
        wg = int(round(w * 1000))
        i = min(range(len(bx)), key=lambda k: load[k])
        if load[i] + wg * q <= cap:
            bx[i][name] = bx[i].get(name, 0) + q
            load[i] += wg * q
            continue
        left = q
        while left > 0:
            room = [k for k in sorted(range(len(bx)), key=lambda k: load[k]) if cap - load[k] >= wg]
            if room:
                k = room[0]
            else:
                bx.append(dict()); load.append(0); k = len(bx) - 1
            take = min(left, (cap - load[k]) // wg)
            bx[k][name] = bx[k].get(name, 0) + take
            load[k] += wg * take
            left -= take
    wmap = {name: w for name, w, _ in items}
    seq = [name for name, _, _ in items]                 # 발주서 순서로 정렬
    out = []
    for b in bx:
        if b:
            out.append([(nm, b[nm], round(b[nm] * wmap[nm], 3)) for nm in seq if nm in b])
    return out


def process(order_src, addr_src):
    """
    반환 dict
      date(date|None), mmdd(str), upload(list[list]), summary(list[list]),
      skipped(list[str]), missing(list[(지점, 어센틱kg, 기타kg)])
    summary 행: [발주서지점, 수취인명, 어센틱kg, 어센틱박스, 기타kg, 기타박스, 총박스, 비고]
    """
    ws = openpyxl.load_workbook(order_src, data_only=True).worksheets[0]

    hrow = next((r for r in range(1, 21)
                 if s(ws.cell(r, 1).value).replace(" ", "") == "품명"), None)
    if hrow is None:
        raise ValueError("발주서에서 '품   명' 헤더행을 찾지 못했습니다. 지점 원두 발주서가 맞는지 확인하세요.")

    # '점포계(EA)' 열 위치로 지점 시작 열을 찾음 (양식마다 열 위치가 달라서 고정값 사용 안 함)
    #   예) 점포계=J → 합계=K, 본점=L  /  점포계=G → 합계=H, 본점=I
    tcol = next((c for c in range(1, ws.max_column + 1)
                 if s(ws.cell(hrow, c).value).replace(" ", "").startswith("점포계")), None)
    if tcol is None:
        raise ValueError("발주서 헤더에서 '점포계(EA)' 열을 찾지 못했습니다.")
    total_col = tcol + 1
    branches = [(c, s(ws.cell(hrow, c).value)) for c in range(tcol + 2, ws.max_column + 1)
                if s(ws.cell(hrow, c).value)]

    auth_rows, other_rows, skipped = [], [], []
    for r in range(hrow + 1, ws.max_row + 1):
        pname = s(ws.cell(r, 1).value)
        if not pname:
            continue
        if pname.replace(" ", "") == AUTHENTIC:
            auth_rows.append(r)
            continue
        m = WEIGHT_RE.search(pname)
        if m:
            w = float(m.group(1)) / (1000 if m.group(2).lower() == "g" else 1)
            other_rows.append((r, w, pname))
        else:
            skipped.append(pname)
    if not auth_rows:
        raise ValueError("발주서에 '어센틱에스프레소블렌드 1kg' 행이 없습니다.")

    addrs = load_addresses(addr_src)

    d = ws.cell(hrow - 2, 1).value
    date = d.date() if isinstance(d, datetime.datetime) else (d if isinstance(d, datetime.date) else None)
    mmdd = date.strftime("%m%d") if date else datetime.date.today().strftime("%m%d")

    auth_name = s(ws.cell(auth_rows[0], 1).value)
    products = [(auth_name, 1.0)] + [(nm, w) for _, w, nm in other_rows]   # 포장표 행 순서
    upload, summary, missing, detail = [], [], [], []
    for c, bname in branches:
        a_kg = round(sum(num(ws.cell(r, c).value) for r in auth_rows), 6)
        o_kg = round(sum(num(ws.cell(r, c).value) * w for r, w, _ in other_rows), 6)
        a_box, o_box = boxes(a_kg), boxes(o_kg)

        # 박스 구성: 어센틱(고르게 분배) → 기타(품목별 배분)  ※ 강릉 지점도 포장표용으로 동일 계산
        a_split = split_even(a_kg, a_box)
        o_items = [(nm, w, int(num(ws.cell(r, c).value))) for r, w, nm in other_rows
                   if num(ws.cell(r, c).value) > 0]
        o_alloc = allocate_other(o_items, o_box) if o_box else []
        o_box = len(o_alloc)
        total = a_box + o_box
        box_list = [("어센틱", [(auth_name, int(q), float(q))], q) for q in a_split] + \
                   [("기타원두", lines, None) for lines in o_alloc]
        boxes_detail = [dict(label=f"({i}/{total})", kind=kind, lines=lines,
                             kg=round(sum(l[2] for l in lines), 3))
                        for i, (kind, lines, _) in enumerate(box_list, 1)]
        qty = {auth_name: a_kg}
        qty.update({nm: q for nm, _, q in o_items})

        if bname in EXCLUDE:
            summary.append([bname, "", a_kg, a_box, o_kg, o_box, 0, "강릉 지점 제외"])
            if total:
                detail.append(dict(branch=bname, recv="", a_kg=a_kg, o_kg=o_kg, local=True,
                                   boxes=boxes_detail, qty=qty))
            continue
        addr, note = match_address(bname, addrs)
        if addr is None:
            if a_kg + o_kg == 0:
                note += " (발주없음)"
            else:
                missing.append((bname, a_kg, o_kg))
            summary.append([bname, "", a_kg, a_box, o_kg, o_box, 0, note])
            continue
        for b, (kind, _, q) in zip(boxes_detail, box_list):
            item = f"{AUTHENTIC_LABEL} {q:g}kg" if kind == "어센틱" else addr[5]
            upload.append([f"{addr[0]}{b['label']}"] + addr[1:5] + [item, addr[6]])
        if total:
            detail.append(dict(branch=bname, recv=addr[0], a_kg=a_kg, o_kg=o_kg, local=False,
                               boxes=boxes_detail, qty=qty))
        summary.append([bname, addr[0], a_kg, a_box, o_kg, o_box, total, "" if total else "발주없음"])

    # 검증: 품목별 지점 합계 vs 발주서 점포계 열
    check = []
    for r in auth_rows + [r for r, _, _ in other_rows]:
        want = ws.cell(r, total_col).value
        if not isinstance(want, (int, float)):
            continue
        got = sum(num(ws.cell(r, c).value) for c, _ in branches)
        if abs(got - want) > 1e-6:
            check.append((s(ws.cell(r, 1).value), want, got))

    return dict(date=date, mmdd=mmdd, upload=upload, summary=summary,
                skipped=skipped, missing=missing, detail=detail, products=products,
                check=check, branch_names=[b for _, b in branches])


def build_workbook(res, src_name=""):
    """결과 dict → xlsx bytes"""
    upload, summary, skipped = res["upload"], res["summary"], res["skipped"]
    wb = openpyxl.Workbook()
    f10 = Font(name="맑은 고딕", size=10)

    # ── 업로드용 (우체국 업로드 양식) ──
    ws = wb.active
    ws.title = "업로드용"
    ws.append(HEADERS)
    for row in upload:
        ws.append([v if v != "" else None for v in row])
    for row in ws.iter_rows():
        for cell in row:
            cell.font = f10
            cell.number_format = "@"
            cell.alignment = Alignment(horizontal="left")
    for col, w in zip("ABCDEFG", [38, 8, 70, 16, 16, 32, 40]):
        ws.column_dimensions[col].width = w

    # ── 박스산출내역 (검증용) ──
    TERRA, LINEN, CREAM, OFF, CHAR, STONE = "8B3A2A", "EDE5DC", "FBF8F5", "F5F0EB", "2C2C2C", "D6CEC8"
    bd = Border(*(Side(style="thin", color=STONE),) * 4)
    ws2 = wb.create_sheet("박스산출내역")
    ws2.sheet_view.showGridLines = False
    ws2.merge_cells("A1:H1")
    ws2["A1"] = f"지점 원두 우체국 박스 산출내역  ({src_name})"
    ws2["A1"].font = Font(name="맑은 고딕", size=14, bold=True, color="FFFFFF")
    ws2["A1"].fill = PatternFill("solid", fgColor=TERRA)
    ws2.row_dimensions[1].height = 26
    ws2.merge_cells("A2:H2")
    ws2["A2"] = "기준: 15kg/박스 올림 · 어센틱에스프레소블렌드 1kg 별도 박스(고르게 분배) · 그 외 중량표기 원두 합산"
    ws2["A2"].font = Font(name="맑은 고딕", size=9, color="6B6056")
    ws2.append([])
    ws2.append(["발주서 지점", "수취인명", "어센틱1kg(kg)", "어센틱 박스", "기타원두(kg)", "기타 박스", "총 박스", "비고"])
    hr = ws2.max_row
    for c in range(1, 9):
        cell = ws2.cell(hr, c)
        cell.font = Font(name="맑은 고딕", size=9, bold=True, color=CHAR)
        cell.fill = PatternFill("solid", fgColor=LINEN)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = bd
    start = hr + 1
    for i, row in enumerate(summary):
        ws2.append(row)
        r = ws2.max_row
        fill = PatternFill("solid", fgColor=CREAM if i % 2 == 0 else OFF)
        if "미등록" in row[7] and row[2] + row[4] > 0:
            fill = PatternFill("solid", fgColor="F0DDD7")
        for c in range(1, 9):
            cell = ws2.cell(r, c)
            bold = c == 7 and bool(row[6])
            cell.font = Font(name="맑은 고딕", size=9, bold=bold, color=TERRA if bold else CHAR)
            cell.fill = fill
            cell.border = bd
            if c in (3, 5):
                cell.number_format = "0.##"
            if 3 <= c <= 7:
                cell.alignment = Alignment(horizontal="right")
    end = ws2.max_row
    ws2.append(["합계 (접수대상)", "",
                f'=SUMIFS(C{start}:C{end},G{start}:G{end},">0")', f'=SUMIFS(D{start}:D{end},G{start}:G{end},">0")',
                f'=SUMIFS(E{start}:E{end},G{start}:G{end},">0")', f'=SUMIFS(F{start}:F{end},G{start}:G{end},">0")',
                f"=SUM(G{start}:G{end})", f"업로드용 행수: {len(upload)}"])
    # 강릉 지점 소계 / 전체 총량(강릉 포함)
    ws2.append(["강릉 지점 소계", "",
                f'=SUMIFS(C{start}:C{end},H{start}:H{end},"강릉 지점 제외")', "",
                f'=SUMIFS(E{start}:E{end},H{start}:H{end},"강릉 지점 제외")', "", "", "우체국 접수 제외"])
    ws2.append(["전체 총량 (강릉 포함)", "",
                f"=SUM(C{start}:C{end})", "", f"=SUM(E{start}:E{end})", "", "", "발주서 전 지점 합계"])
    for r in range(ws2.max_row - 2, ws2.max_row + 1):
        for c in range(1, 9):
            cell = ws2.cell(r, c)
            cell.font = Font(name="맑은 고딕", size=10, bold=True, color=TERRA)
            cell.fill = PatternFill("solid", fgColor=LINEN)
            cell.border = bd
            if c in (3, 5):
                cell.number_format = "0.##"
    if skipped:
        ws2.append([])
        ws2.append(["※ 중량표기 없어 제외된 품목: " + ", ".join(skipped)])
        ws2.cell(ws2.max_row, 1).font = Font(name="맑은 고딕", size=8, color="6B6056")
    for col, w in zip("ABCDEFGH", [14, 36, 13, 11, 13, 10, 9, 30]):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = f"A{start}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ───────────────────────── 포장표 (박스 배분용) ─────────────────────────
def build_packing(res, src_name=""):
    """포장표.xlsx bytes: ① 지점별 포장표(발주서 형태) ② 박스별 구성"""
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.pagebreak import Break
    TERRA, MID, LINEN, CREAM, OFF, CHAR, GRAY, STONE, PEAK = \
        "8B3A2A", "C4644A", "EDE5DC", "FBF8F5", "F5F0EB", "2C2C2C", "6B6056", "D6CEC8", "F5E8D8"
    thin = Side(style="thin", color=STONE)
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    F = lambda sz=9, b=False, c=CHAR: Font(name="맑은 고딕", size=sz, bold=b, color=c)
    fill = lambda c: PatternFill("solid", fgColor=c)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    det = res["detail"]
    date_txt = res["date"].strftime("%Y-%m-%d") if res["date"] else res["mmdd"]

    wb = openpyxl.Workbook()

    # ── ① 포장표 : 행=원두, 열=지점 (인쇄용으로 지점 14곳씩 블록 분할) ──
    ws = wb.active
    ws.title = "포장표"
    ws.sheet_view.showGridLines = False
    used = [(nm, w) for nm, w in res["products"] if any(d["qty"].get(nm, 0) for d in det)]
    auth_nm = res["products"][0][0]
    nblk = max(1, math.ceil(len(det) / 14))
    PER = max(1, math.ceil(len(det) / nblk))          # 블록당 지점 수(최대 14, 고르게)
    width = 2 + PER + 2
    loc = [d for d in det if d.get("local")]
    shp = [d for d in det if not d.get("local")]
    loc_box = sum(len(d["boxes"]) for d in loc)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    tot_a = sum(d["a_kg"] for d in det); tot_o = sum(d["o_kg"] for d in det)
    ws["A1"] = (f"지점 원두 포장표  ·  발주일 {date_txt}  ·  우체국 {len(shp)}곳 {len(res['upload'])}박스"
                f" + 강릉 {len(loc)}곳 {loc_box}박스  ·  어센틱 {tot_a:g}kg / 기타 {round(tot_o, 3):g}kg")
    ws["A1"].font = F(13, True, "FFFFFF"); ws["A1"].fill = fill(TERRA)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws["A2"] = ("숫자 = 수량(EA)  ·  (강릉) = 강릉권 지점(우체국 접수 제외)  ·  주소 미등록(미오픈) 지점 제외  ·  "
                "박스 구성: 어 = 어센틱 박스별 kg, 기 = 기타 원두 박스별 kg  ·  품목별 상세는 '박스별 구성' 시트")
    ws["A2"].font = F(8, False, GRAY)

    def write_block(top, chunk):
        nb = len(chunk)
        sub1, sub2 = 3 + nb, 4 + nb                # 소계(EA), 소계(kg)
        heads = ["품   명", "규격"] + [d["branch"] + ("\n(강릉)" if d.get("local") else "") for d in chunk] \
                + ["소계\n(EA)", "소계\n(kg)"]
        for c, h in enumerate(heads, 1):
            cell = ws.cell(top, c, h)
            is_loc = 3 <= c < 3 + nb and chunk[c - 3].get("local")
            cell.font = F(9, True, TERRA if is_loc else CHAR)
            cell.fill = fill("F0DDD7" if is_loc else LINEN); cell.alignment = center; cell.border = bd
        ws.row_dimensions[top].height = 30
        r0 = top + 1
        for i, (nm, w) in enumerate(used):
            r = r0 + i
            is_auth = nm == auth_nm
            rf = fill(PEAK) if is_auth else fill(CREAM if i % 2 == 0 else OFF)
            ws.cell(r, 1, nm); ws.cell(r, 2, w)
            for j, d in enumerate(chunk):
                q = d["qty"].get(nm, 0)
                ws.cell(r, 3 + j, q if q else None)
            ws.cell(r, sub1, f"=SUM({L(3)}{r}:{L(2 + nb)}{r})")
            ws.cell(r, sub2, f"={L(sub1)}{r}*B{r}")
            for c in range(1, sub2 + 1):
                cell = ws.cell(r, c)
                cell.border = bd; cell.fill = rf
                cell.font = F(9, is_auth or c >= sub1)
                cell.alignment = Alignment(horizontal=None if c == 1 else "center", vertical="center")
                if c == 2:
                    cell.number_format = '0.##"kg"'
                if c >= sub1:
                    cell.number_format = "0.##"
        rA, rN = r0, r0 + len(used) - 1
        srow = rN + 1
        labels = ["어센틱 1kg (kg)", "기타 원두 (kg)", "총 중량 (kg)", "박스 수", "박스 구성"]
        for k, lab in enumerate(labels):
            r = srow + k
            ws.cell(r, 1, lab)
            for c in range(3, sub1 + 1):
                col, j = L(c), c - 3
                if k == 0:
                    v = f"={col}{rA}"
                elif k == 1:
                    v = f"=SUMPRODUCT({col}{rA + 1}:{col}{rN},$B${rA + 1}:$B${rN})" if rN > rA else 0
                elif k == 2:
                    v = f"={col}{srow}+{col}{srow + 1}"
                elif k == 3:
                    v = len(chunk[j]["boxes"]) if j < nb else f"=SUM({L(3)}{r}:{L(2 + nb)}{r})"
                else:
                    if j < nb:
                        bxs = chunk[j]["boxes"]
                        a = [f"{b['kg']:g}" for b in bxs if b["kind"] == "어센틱"]
                        o = [f"{b['kg']:g}" for b in bxs if b["kind"] != "어센틱"]
                        v = "\n".join(x for x in [("어" + "·".join(a)) if a else "",
                                                   ("기" + "·".join(o)) if o else ""] if x)
                    else:
                        v = None
                ws.cell(r, c, v)
            for c in range(1, sub2 + 1):
                cell = ws.cell(r, c)
                cell.border = bd; cell.fill = fill(LINEN)
                cell.font = F(9 if k < 4 else 8, k in (2, 3), TERRA if k in (2, 3) else CHAR)
                cell.alignment = Alignment(horizontal="center" if c > 1 else None, vertical="center", wrap_text=True)
                if k < 3 and c > 2:
                    cell.number_format = "0.##"
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            if k == 4:
                est = [sum(math.ceil(max(len(seg), 1) / 8) for seg in str(ws.cell(r, c).value or "").split("\n"))
                       for c in range(3, sub1)]
                ws.row_dimensions[r].height = max(30, 12 * max(est + [1]) + 6)
        return srow + 5

    top = 4
    chunks = [det[i:i + PER] for i in range(0, len(det), PER)] or [[]]
    for bi, chunk in enumerate(chunks):
        end = write_block(top, chunk)
        if bi < len(chunks) - 1:
            ws.row_breaks.append(Break(id=end))       # 블록마다 인쇄 페이지 나눔
        top = end + 2

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 6.5
    for c in range(3, 3 + PER):
        ws.column_dimensions[L(c)].width = 7.5
    ws.column_dimensions[L(3 + PER)].width = 7
    ws.column_dimensions[L(4 + PER)].width = 7.5
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4

    # ── ①-2 원두별 분배 : 원두(품명) 하나씩 집어서 지점 박스에 뿌리는 순서 ──
    #   A 지점 · B 수량 · C 확인  |  D 간격  |  E 박스 · F 구분 · G 중량
    wd = wb.create_sheet("원두별 분배")
    wd.sheet_view.showGridLines = False
    wd.merge_cells("A1:G1")
    wd["A1"] = f"원두별 분배표  ·  발주일 {date_txt}"
    wd["A1"].font = F(14, True, "FFFFFF"); wd["A1"].fill = fill(TERRA)
    wd["A1"].alignment = Alignment(vertical="center", indent=1)
    wd.row_dimensions[1].height = 28
    wd.merge_cells("A2:G2")
    wd["A2"] = ("원두 하나씩 꺼내 지점 박스에 나눠 담기  ·  원두 1종 = 1블록  ·  지점 순서는 발주서와 같음  ·  "
                "(강릉) = 강릉권 지점  ·  박스 = 송장 수취인명 뒤 번호  ·  담은 뒤 ☐ 체크")
    wd["A2"].font = F(8, False, GRAY)
    LOCFILL = "F0DDD7"
    COLS = [1, 2, 3, 5, 6, 7]                  # D(4)열은 간격
    per = {}
    for d in det:
        for b in d["boxes"]:
            for nm, q, kg in b["lines"]:
                per.setdefault(nm, []).append((d, b["label"], q, kg))
    r = 4
    ROWS_PER_PAGE = 48                          # A4 세로, 폭 맞춤 기준 대략 행 수
    page_used = 3
    for nm, _w in res["products"]:
        items = per.get(nm)
        if not items:
            continue
        need = len(items) + 3
        if page_used + need > ROWS_PER_PAGE and page_used > 3:
            wd.row_breaks.append(Break(id=r - 1))
            page_used = 0
        tot_q = sum(x[2] for x in items)
        tot_kg = round(sum(x[3] for x in items), 3)
        n_br = len({x[0]["branch"] for x in items})
        # 원두 제목줄: A~C 원두명 / E 총수량 / F 지점수 / G 총중량
        wd.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        wd.cell(r, 1, nm)
        wd.cell(r, 5, f"총 {tot_q:g}개")
        wd.cell(r, 6, f"{n_br}곳")
        wd.cell(r, 7, f"{tot_kg:g}kg")
        bar = TERRA if nm == auth_nm else MID
        for c in COLS:
            cell = wd.cell(r, c)
            cell.fill = fill(bar)
            cell.font = F(13 if c == 1 else 10, True, "FFFFFF")
            cell.alignment = Alignment(horizontal="left" if c == 1 else "center", vertical="center",
                                       indent=1 if c == 1 else 0)
        wd.row_dimensions[r].height = 24
        r += 1
        for c, h in zip(COLS, ["지점", "수량(EA)", "확인", "박스", "구분", "중량(kg)"]):
            cell = wd.cell(r, c, h)
            cell.font = F(9, True); cell.fill = fill(LINEN); cell.alignment = center; cell.border = bd
        r += 1
        prev = None
        for i, (d, label, q, kg) in enumerate(items):
            loc = d.get("local")
            same = prev == d["branch"]
            vals = {1: "" if same else d["branch"] + (" (강릉)" if loc else ""), 2: q, 3: "☐",
                    5: label, 6: "강릉" if loc else "우체국", 7: kg}
            rf = fill(LOCFILL if loc else (CREAM if i % 2 == 0 else OFF))
            for c in COLS:
                cell = wd.cell(r, c, vals[c])
                cell.fill = rf; cell.border = bd
                cell.alignment = Alignment(horizontal="left" if c == 1 else "center", vertical="center",
                                           indent=1 if c == 1 else 0)
                if c == 2:
                    cell.font = F(13, True, TERRA)
                elif c == 3:
                    cell.font = F(12, False, CHAR)
                elif c == 1:
                    cell.font = F(11, True, CHAR)
                else:
                    cell.font = F(9, False, GRAY)
                if c == 7:
                    cell.number_format = "0.##"
            wd.row_dimensions[r].height = 19
            prev = d["branch"]
            r += 1
        wd.cell(r, 1, "합계"); wd.cell(r, 2, tot_q); wd.cell(r, 7, tot_kg)
        for c in COLS:
            cell = wd.cell(r, c)
            cell.fill = fill(LINEN); cell.border = bd
            cell.font = F(10, True, TERRA)
            cell.alignment = Alignment(horizontal="left" if c == 1 else "center", vertical="center",
                                       indent=1 if c == 1 else 0)
            if c == 7:
                cell.number_format = "0.##"
        r += 2
        page_used += need + 1
    for col, w in zip("ABCDEFG", [24, 11, 7, 2, 9, 8, 10]):
        wd.column_dimensions[col].width = w
    wd.freeze_panes = "A4"
    wd.page_setup.orientation = "portrait"
    wd.page_setup.paperSize = wd.PAPERSIZE_A4
    wd.page_setup.fitToWidth = 1
    wd.page_setup.fitToHeight = 0
    wd.sheet_properties.pageSetUpPr.fitToPage = True
    wd.page_margins.top = wd.page_margins.bottom = 0.4

    # ── ② 박스별 구성 ──
    ws2 = wb.create_sheet("박스별 구성")
    ws2.sheet_view.showGridLines = False
    ws2.merge_cells("A1:H1")
    ws2["A1"] = f"박스별 구성  ·  발주일 {date_txt}  ·  우체국 {len(res['upload'])}박스 + 강릉 {loc_box}박스"
    ws2["A1"].font = F(14, True, "FFFFFF"); ws2["A1"].fill = fill(TERRA)
    ws2["A1"].alignment = Alignment(vertical="center", indent=1)
    ws2.row_dimensions[1].height = 28
    hd = ["지점", "박스", "구분", "품명", "수량(EA)", "중량(kg)", "박스 중량(kg)", "확인"]
    ws2.append([]); ws2.append(hd)
    for c in range(1, 9):
        cell = ws2.cell(3, c)
        cell.font = F(9, True); cell.fill = fill(LINEN); cell.alignment = center; cell.border = bd
    r = 4
    for bi, d in enumerate(det):
        for k, b in enumerate(d["boxes"]):
            bfill = fill(CREAM if k % 2 == 0 else OFF)
            for li, (nm, q, kg) in enumerate(b["lines"]):
                vals = [(d["branch"] + ("(강릉)" if d.get("local") else "")) if (k == 0 and li == 0) else None,
                        b["label"] if li == 0 else None,
                        b["kind"] if li == 0 else None,
                        nm, q, kg,
                        b["kg"] if li == 0 else None,
                        "☐" if li == 0 else None]
                for c, v in enumerate(vals, 1):
                    cell = ws2.cell(r, c, v)
                    cell.fill = bfill
                    cell.font = F(9, c in (1, 2, 7) and v is not None, TERRA if c == 7 else CHAR)
                    cell.alignment = Alignment(horizontal="left" if c == 4 else "center", vertical="center")
                    top = Side(style="medium", color=TERRA) if (k == 0 and li == 0) else \
                          (Side(style="thin", color=GRAY) if li == 0 else thin)
                    cell.border = Border(left=thin, right=thin, top=top, bottom=thin)
                    if c in (6, 7):
                        cell.number_format = "0.##"
                r += 1
    ws2.cell(r + 1, 1, "※ 박스마다 15kg 이하 · 같은 원두는 가능한 한 한 박스에 담도록 배분").font = F(8, False, GRAY)
    for col, w in zip("ABCDEFGH", [15, 7, 9, 34, 9, 9, 12, 6]):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A4"
    ws2.page_setup.orientation = "portrait"
    ws2.page_setup.paperSize = ws2.PAPERSIZE_A4
    ws2.page_setup.fitToWidth = 1
    ws2.page_setup.fitToHeight = 0
    ws2.sheet_properties.pageSetUpPr.fitToPage = True
    ws2.print_title_rows = "3:3"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
